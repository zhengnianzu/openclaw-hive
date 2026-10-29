# -*- coding: utf-8 -*-
"""任务详情导出：脱敏 + 收集数据 + 组装 Excel 工作簿。

导出的口径来自两处既有实现，必须保持一致：
  - 关键项 config  ← api.services.create_params.build_create_params
    （即详情页「配置信息 → 创建参数」卡片）
  - 成功/失败数量  ← instances._sync_instance_status 数 complete.jsonl / failed.jsonl
    （即 Dashboard 的「成功」「失败」列）

安全：API Key / token 等敏感值**只在这里脱敏**，不依赖前端展示层。
关键项 Sheet 里的 key 走 mask_secret，掩码而非明文。
"""
from __future__ import annotations

from io import BytesIO
from urllib.parse import quote

from .create_params import build_create_params

# 任务明细行数上限：超过就只给提示行，避免内存与文件体积失控
_TASK_RECORD_ROW_LIMIT = 20000

MASK_PLACEHOLDER = "****"

# 命中即视为敏感值。按小写比较，用子串匹配（api_key / apiKey / ANTHROPIC_AUTH_TOKEN 都能覆盖）
_SECRET_KEY_WORDS = (
    "api_key", "apikey", "api-key",
    "secret", "access_key", "access-key",
    "password", "passwd", "private_key",
    "token", "auth_token", "access_token",
)

def mask_secret(value) -> str:
    """API Key 等敏感值的统一掩码。

    与前端 maskKey（InstanceDetail.vue）的展示效果一致，但更严格：
    前端对长度 <= 8 的值原样返回，这里一律整体打码，不泄漏短 key。
    """
    if value is None:
        return ""
    s = str(value).strip()
    if not s:
        return ""
    if s == MASK_PLACEHOLDER:
        return s
    if len(s) <= 8:
        return MASK_PLACEHOLDER
    return s[:4] + "***" + s[-4:]


def _fmt_time(v) -> str:
    if not v:
        return ""
    return str(v).replace("T", " ")[:19]


def _rate(completed: int, failed: int) -> str:
    done = (completed or 0) + (failed or 0)
    if not done:
        return "—"
    return f"{completed / done * 100:.1f}%"


def collect_rows(instance_ids: list[str], conn, sync_status=None, with_resources=None) -> list[dict]:
    """按给定顺序收集实例行（含 create_params / 最新计数）。

    - 去重保序；
    - 查不到的 id 保留占位行（missing=True），而不是静默丢弃 —— 用户选了 10 个
      只导出 9 个会困惑；
    - sync_status 用于把 completed/failed 刷成最新值；传入的必须是**只读**版本
      （instances._sync_instance_status 默认 persist=False），读路径写库会撞锁；
    - with_resources 用于补「单沙箱资源」列（instances._with_resources，带进程内缓存）。
    """
    seen = set()
    ordered = []
    for i in instance_ids:
        if not i or i in seen:
            continue
        seen.add(i)
        ordered.append(i)

    if not ordered:
        return []

    placeholders = ",".join("?" * len(ordered))
    rows = conn.execute(
        f"SELECT * FROM task_instances WHERE id IN ({placeholders})", tuple(ordered)
    ).fetchall()
    by_id = {r["id"]: dict(r) for r in rows}

    out = []
    for i in ordered:
        row = by_id.get(i)
        if row is None:
            out.append({"id": i, "missing": True})
            continue
        if sync_status is not None:
            try:
                row = sync_status(row)
            except Exception:
                pass  # 计数刷新失败不影响导出其余内容
        if with_resources is not None:
            try:
                row = with_resources(row)
            except Exception:
                pass  # 快照缺失/损坏时资源列为空即可
        params, meta = build_create_params(row)
        row["_params"] = params
        row["_params_meta"] = meta
        out.append(row)
    return out


# 漏斗档位顺序：与 api/routers/batch_output.py 的 _TRAJ_LEVELS 必须一致。
# 这里手写常量而不是 import，避免 services 反向依赖 routers。
_FUNNEL_LEVELS = ["L0", "L1", "L1.5", "L2", "L3"]


def _empty_funnel() -> dict:
    d = {lv: 0 for lv in _FUNNEL_LEVELS}
    d.update({"graded": 0, "unevaluated": 0, "failed": 0, "total": 0, "task_done": 0})
    return d


def collect_funnel_counts(conn, instance_ids: list[str]) -> dict[str, dict]:
    """批量取各实例的「会话漏斗」计数（task_records.traj_level 聚合）。

    口径逐字对齐 api/routers/batch_output.py 的 get_shallow，保证导出数字与详情页
    「会话漏斗」卡片一致：
      graded      = traj_level IN ('L0','L1','L1.5','L2','L3')
      unevaluated = traj_level IS NULL（未被浅层分析覆盖 / 仍在排队）
      failed      = traj_level = 'failed'
      total       = 该实例全部行
      task_done   = LEFT JOIN task_traj_records 后的 SUM(task_done)

    注意 total 与 graded+unevaluated+failed 不一定相等：库里还存在 traj_level='none'
    的残差行（采集侧过期 stats 的形态），三者都不算。

    一条 SQL 出全部实例，走 idx_task_records_inst。不要复用
    /instances/{id}/shallow/export 的实现——那是把整实例所有行拉进 Python 再数，
    批量导出会变成 N 次全表扫。
    """
    ids = [i for i in dict.fromkeys(instance_ids) if i]
    out = {i: _empty_funnel() for i in ids}
    if not ids:
        return out

    levels = set(_FUNNEL_LEVELS)
    placeholders = ",".join("?" * len(ids))
    rows = conn.execute(
        f"""SELECT tr.instance_id AS iid, tr.traj_level AS lvl, COUNT(*) AS n,
                   SUM(CASE WHEN ttr.task_done THEN 1 ELSE 0 END) AS done_n
            FROM task_records tr
            LEFT JOIN task_traj_records ttr
              ON ttr.instance_id = tr.instance_id
             AND ttr.config_name = tr.config_name
            WHERE tr.instance_id IN ({placeholders})
            GROUP BY tr.instance_id, tr.traj_level""",
        tuple(ids),
    ).fetchall()

    for r in rows:
        d = out[r["iid"]]
        n = r["n"] or 0
        lvl = r["lvl"]
        d["total"] += n
        d["task_done"] += r["done_n"] or 0
        if lvl is None:
            d["unevaluated"] += n
        elif lvl == "failed":
            d["failed"] += n
        elif lvl in levels:
            d[lvl] += n
            d["graded"] += n
        # 其余（如 'none'）只计入 total，不进任何桶——与 get_shallow 一致
    return out


def _fit_columns(ws):
    """列宽按内容最大宽 + 6（同 batch_output.export_shallow_funnel）。"""
    for col_idx, cell in enumerate(ws[1], start=1):
        max_len = len(str(cell.value or ""))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            for c in row:
                max_len = max(max_len, len(str(c.value or "")))
        ws.column_dimensions[cell.column_letter].width = min(max_len + 6, 120)


def _write_header(ws, headers):
    from openpyxl.styles import Font, PatternFill

    ws.append(headers)
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    ws.freeze_panes = "A2"


def _obs_path(dir_path: str, bucket: str) -> str:
    """OBS 相对目录拼上桶名；已是绝对路径或缺桶名则原样返回。

    逻辑同前端 InstanceDetail.vue 的 obsPath()，保证与详情页显示一致。
    """
    if not dir_path:
        return ""
    d = str(dir_path)
    if "://" in d or d.startswith("/") or not bucket:
        return d
    b = str(bucket).rstrip("/")
    return f"{b}/{d.lstrip('/')}"


def _sandbox_res(strs) -> str:
    """单沙箱资源展示：cpu_request/cpu_limit 核 · memory_request/memory_limit GiB。"""
    def g(k):
        v = strs.get(k)
        return "" if v in (None, "") else v

    cpu_req, cpu_lim = g("cpu_request"), g("cpu_limit")
    mem_req, mem_lim = g("memory_request"), g("memory_limit")
    if not any([cpu_req, cpu_lim, mem_req, mem_lim]):
        return ""
    return f"{cpu_req}/{cpu_lim} 核 · {mem_req}/{mem_lim} GiB"


def _harness_label(params: dict) -> str:
    name = params.get("harness_config_name")
    if name:
        ver = params.get("harness_config_version")
        return f"{name} / v{ver}" if ver else name
    return params.get("harness_config_label") or params.get("harness_config_id") or ""


# ---------------------------------------------------------------------------
# Sheet 组装
# ---------------------------------------------------------------------------

def _sheet_overview(wb, rows, funnel=None, title="任务总览"):
    ws = wb.create_sheet(title)
    funnel = funnel or {}
    funnel_headers = (
        [f"会话{lv}" for lv in _FUNNEL_LEVELS]
        + ["已分级", "未评估", "分级失败", "会话总数", "完成标记"]
    )
    headers = [
        "任务名称", "实例ID", "所属队列", "队列位置", "Harness", "状态",
        "创建者", "创建时间", "开始时间", "结束时间", "并发数",
        "任务总数", "成功数", "失败数", "未完成数", "成功率",
        "失败原因", "关键项来源",
    ] + funnel_headers
    _write_header(ws, headers)
    tot_total = tot_ok = tot_fail = 0
    tot_funnel = _empty_funnel()
    for r in rows:
        f = funnel.get(r.get("id")) or _empty_funnel()
        funnel_cells = [f[lv] for lv in _FUNNEL_LEVELS] + [
            f["graded"], f["unevaluated"], f["failed"], f["total"], f["task_done"],
        ]
        if r.get("missing"):
            ws.append([f"（任务不存在：{r['id']}）", r["id"]]
                      + [""] * (len(headers) - 2))
            continue
        total = r.get("total_tasks") or 0
        ok = r.get("completed_tasks") or 0
        fail = r.get("failed_tasks") or 0
        tot_total += total
        tot_ok += ok
        tot_fail += fail
        for k in tot_funnel:
            tot_funnel[k] += f[k]
        meta = r.get("_params_meta") or {}
        has_params = bool(r.get("_params"))
        if not has_params:
            source = "无 create_params"
        else:
            source = "快照回填" if meta.get("has_snapshot") else "仅 create_params"
        if meta.get("snapshot_error"):
            source += f"（快照解析失败：{meta['snapshot_error']}）"
        ws.append([
            r.get("name") or "",
            r.get("id") or "",
            r.get("queue_id") or "未入队",
            r.get("queue_position") or 0,
            r.get("harness_type") or "",
            r.get("status") or "",
            r.get("created_by") or "",
            _fmt_time(r.get("created_at")),
            _fmt_time(r.get("started_at")),
            _fmt_time(r.get("stopped_at")),
            r.get("concurrent_num") or 0,
            total,
            ok,
            fail,
            # total_tasks 是启动前估算的，实际 complete+failed 可能超出它（重跑、补跑），
            # 直接相减会出现负数，钳到 0。
            max(0, total - ok - fail) if total else "",
            _rate(ok, fail),
            r.get("error_summary") or "",
            source,
        ] + funnel_cells)
    if len(rows) > 1:
        done = tot_ok + tot_fail
        ws.append([
            "合计", "", "", "", "", "", "", "", "", "", "",
            tot_total, tot_ok, tot_fail,
            max(0, tot_total - done) if tot_total else "",
            f"{tot_ok / done * 100:.1f}%" if done else "—",
            "", "",
        ] + [tot_funnel[lv] for lv in _FUNNEL_LEVELS] + [
            tot_funnel["graded"], tot_funnel["unevaluated"], tot_funnel["failed"],
            tot_funnel["total"], tot_funnel["task_done"],
        ])
    _fit_columns(ws)
    return ws


def _sheet_key_params(wb, rows, title="创建参数-关键项"):
    """关键项：列顺序严格对齐详情页「创建参数」卡片的五个分组。"""
    ws = wb.create_sheet(title)
    # agents 列数按本批任务里 agents 的最大长度动态生成，保持表结构稳定
    max_agents = 0
    for r in rows:
        agents = (r.get("_params") or {}).get("agents") or []
        if isinstance(agents, list):
            max_agents = max(max_agents, len(agents))
    max_agents = max(max_agents, 1)

    headers = [
        "任务名称", "实例ID",
        # 基本配置
        "基本.任务标识", "基本.Harness类型", "基本.Invite Code", "基本.Harness配置",
        # OBS配置
        "OBS.技能目录", "OBS.默认技能目录", "OBS.默认技能", "OBS.Agent目录",
        "OBS.用户Config目录", "OBS.用户Profile目录", "OBS.轨迹保存路径",
        # 模型配置
        "模型.Base URL", "模型.API Key", "模型.API类型", "模型.模型ID",
    ]
    for i in range(1, max_agents + 1):
        headers += [
            f"Agent{i}.名称", f"Agent{i}.Provider", f"Agent{i}.Base URL",
            f"Agent{i}.API Key", f"Agent{i}.API类型", f"Agent{i}.模型ID",
        ]
    # 无 agents 时的 user_proxy 回退列（与详情页「用户模拟配置」的另一分支对应）
    headers += ["用户模拟.模型名称", "用户模拟.Base URL", "用户模拟.API Key"]
    headers += [
        # 高级配置
        "高级.起始索引", "高级.任务总数", "高级.沙箱配置", "高级.沙箱标识",
        "高级.沙箱块", "高级.沙箱文件", "高级.Key沙箱后缀", "高级.资源池",
        "高级.单沙箱资源", "高级.镜像名称", "高级.代码仓ID",
    ]
    _write_header(ws, headers)

    for r in rows:
        if r.get("missing"):
            ws.append([f"（任务不存在：{r['id']}）", r["id"]] + [""] * (len(headers) - 2))
            continue
        p = r.get("_params") or {}
        if not p:
            ws.append([r.get("name") or "", r.get("id") or ""] +
                      ["—"] * (len(headers) - 2))
            continue
        bucket = p.get("obs_bucket") or ""
        row = [
            r.get("name") or "", r.get("id") or "",
            p.get("task_name") or "",
            p.get("harness_type") or "",
            p.get("invite_code") or "",
            _harness_label(p),
            _obs_path(p.get("skill_dir"), bucket),
            _obs_path(p.get("default_skill_dir"), bucket),
            p.get("default_skills") or "",
            _obs_path(p.get("agent_dir"), bucket),
            _obs_path(p.get("user_config_dir"), bucket),
            _obs_path(p.get("user_profile_dir"), bucket),
            _obs_path(p.get("traj_save_path"), p.get("traj_save_bucket") or bucket),
            p.get("model_base_url") or "",
            mask_secret(p.get("model_api_key")),
            p.get("model_api_type") or "",
            p.get("model_id") or "",
        ]
        agents = p.get("agents") if isinstance(p.get("agents"), list) else []
        for i in range(max_agents):
            ag = agents[i] if i < len(agents) and isinstance(agents[i], dict) else {}
            row += [
                ag.get("name") or "",
                ag.get("provider") or "",
                ag.get("base_url") or "",
                mask_secret(ag.get("api_key")),
                ag.get("api") or "",
                ag.get("model") or "",
            ]
        row += [
            p.get("user_proxy_model_name") or "",
            p.get("user_proxy_base_url") or "",
            mask_secret(p.get("user_proxy_api_key")),
        ]
        total_num = p.get("total_num")
        row += [
            p.get("start_index", 0) if p.get("start_index") is not None else 0,
            total_num if total_num else "不限",
            p.get("sandbox_config_name") or "",
            p.get("task_name") or "",
            p.get("sandbox_active_key") or "",
            p.get("sandbox_file") or "",
            "已开启" if p.get("suffix_api_keys") else "未开启",
            p.get("sandbox_namespace") or r.get("sandbox_namespace") or "",
            _sandbox_res(r),
            p.get("sandbox_image") or p.get("image_name") or "",
            p.get("code_repo_id") or "",
        ]
        ws.append(row)
    _fit_columns(ws)
    return ws


def _sheet_task_records(wb, conn, instance_ids, title="任务明细"):
    """per-task 明细（task_records）。没有数据就返回 None，不建空 Sheet。"""
    ids = [i for i in instance_ids if i]
    if not ids:
        return None
    placeholders = ",".join("?" * len(ids))
    total = conn.execute(
        f"SELECT COUNT(*) AS c FROM task_records WHERE instance_id IN ({placeholders})",
        tuple(ids),
    ).fetchone()["c"]
    if not total:
        return None

    ws = wb.create_sheet(title)
    _write_header(ws, [
        "实例ID", "任务名称", "task_idx", "config_name", "状态",
        "错误码", "错误分类", "评估分", "completion", "gate", "轨迹层级",
    ])
    if total > _TASK_RECORD_ROW_LIMIT:
        ws.append([f"（明细共 {total} 行，超过 {_TASK_RECORD_ROW_LIMIT} 行已省略，请按单任务导出）"] +
                  [""] * 10)
        _fit_columns(ws)
        return ws

    rows = conn.execute(
        f"""SELECT tr.instance_id, ti.name AS inst_name, tr.task_idx, tr.config_name,
                   tr.status, tr.error_code, tr.error_category,
                   tr.eval_score, tr.eval_completion, tr.gate, tr.traj_level
            FROM task_records tr
            LEFT JOIN task_instances ti ON ti.id = tr.instance_id
            WHERE tr.instance_id IN ({placeholders})
            ORDER BY tr.instance_id, tr.task_idx""",
        tuple(ids),
    ).fetchall()
    for r in rows:
        ws.append([
            r["instance_id"], r["inst_name"] or "", r["task_idx"] or "",
            r["config_name"] or "", r["status"] or "",
            r["error_code"] or "", r["error_category"] or "",
            r["eval_score"], r["eval_completion"], r["gate"], r["traj_level"] or "",
        ])
    _fit_columns(ws)
    return ws


def _sheet_queue_info(wb, queue: dict, effective_concurrent: int, task_count: int,
                      title="队列信息"):
    ws = wb.create_sheet(title)
    _write_header(ws, ["项目", "值"])
    for k, v in [
        ("队列ID", queue.get("id")),
        ("队列名称", queue.get("name") or queue.get("id")),
        ("并发上限（队列设置）", queue.get("max_concurrent")),
        ("有效并发（受全局上限约束）", effective_concurrent),
        ("长尾跳过", "开启" if queue.get("tail_skip_enabled") else "关闭"),
        ("长尾阈值", f"{queue.get('tail_skip_percent') or 10}%"),
        ("描述", queue.get("description") or ""),
        ("创建者", queue.get("created_by") or ""),
        ("创建时间", _fmt_time(queue.get("created_at"))),
        ("任务总数", task_count),
    ]:
        ws.append([k, "" if v is None else v])
    _fit_columns(ws)
    return ws


def build_workbook(rows, conn=None, queue=None, effective_concurrent=None,
                   task_count=None, include_task_records=True,
                   all_queues=None, truncated_note=None):
    """组装导出工作簿。

    rows                   —— collect_rows() 的结果
    queue                  —— 队列入口时传 task_queues 行，会在最前面插「队列信息」
    include_task_records   —— 是否附加 per-task 明细 Sheet
    all_queues             —— (队列行列表, collect_queue_stats 结果)，传入时插「队列汇总」Sheet
    truncated_note         —— 「导出全部」被截断时的提示，写进概览页
    """
    from openpyxl import Workbook

    wb = Workbook()
    # Workbook() 自带一个空 Sheet，改造成概览页而不是留下空页
    wb.active.title = "导出概览"

    overview = wb["导出概览"]
    _write_header(overview, ["项目", "值"])
    valid = [r for r in rows if not r.get("missing")]
    missing = [r for r in rows if r.get("missing")]
    tot_total = sum(r.get("total_tasks") or 0 for r in valid)
    tot_ok = sum(r.get("completed_tasks") or 0 for r in valid)
    tot_fail = sum(r.get("failed_tasks") or 0 for r in valid)
    # 会话漏斗计数：一条聚合 SQL 出全部实例；conn 为 None（单元测试）时给空值，列仍在
    funnel = collect_funnel_counts(conn, [r.get("id") for r in valid]) if conn is not None else {}
    from datetime import datetime as _dt
    if queue:
        scope = f"队列 {queue.get('id')}"
    elif all_queues:
        scope = f"全部数据（{len(all_queues[0])} 个队列）"
    else:
        scope = f"选中 {len(rows)} 个任务"
    info = [
        ("导出时间", _dt.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("导出范围", scope),
        ("任务数", f"{len(valid)}（另有 {len(missing)} 个任务不存在）" if missing else len(valid)),
        ("任务总数", tot_total),
        ("成功数合计", tot_ok),
        ("失败数合计", tot_fail),
        ("整体成功率", f"{tot_ok / (tot_ok + tot_fail) * 100:.1f}%" if (tot_ok + tot_fail) else "—"),
        ("数据口径", "成功/失败数取自各任务 output 目录的 complete.jsonl / failed.jsonl；"
                     "成功率 = 成功 ÷ (成功 + 失败)，任务总数为启动前的估算值，仅供参考"),
        ("会话漏斗口径", "取自 task_records.traj_level（与详情页「会话漏斗」卡片同源）；"
                        "已分级 = L0-L3 合计，未评估 = 尚未分析/排队中，分级失败单独计数。"
                        "会话总数是实际会话行数，与上面「任务总数」（启动前估算值）口径不同，"
                        "两者本就不相等；会话总数可能略大于三者之和（残差为 traj_level='none' 的无效分级记录）"),
        ("安全说明", "API Key / token 等敏感值已脱敏（掩码），本文件可对外流转"),
    ]
    if truncated_note:
        info.insert(3, ("注意", truncated_note))
    for k, v in info:
        overview.append([k, v])
    _fit_columns(overview)

    if queue:
        _sheet_queue_info(wb, queue, effective_concurrent, task_count)

    if all_queues:
        _sheet_queue_summary(wb, all_queues[0], all_queues[1])

    _sheet_overview(wb, rows, funnel=funnel)
    _sheet_key_params(wb, rows)

    if include_task_records and conn is not None:
        _sheet_task_records(wb, conn, [r.get("id") for r in valid])

    return wb


def collect_queue_stats(conn, queue_ids: list[str]) -> dict:
    """按队列聚合任务状态与成功/失败计数（导出队列列表用）。

    一条 SQL 出全部队列的统计，避免逐队列查询。返回 {queue_id: {...}}。
    """
    if not queue_ids:
        return {}
    placeholders = ",".join("?" * len(queue_ids))
    rows = conn.execute(
        f"""SELECT queue_id,
                   COUNT(*) AS task_count,
                   SUM(CASE WHEN status IN ('preparing','running') THEN 1 ELSE 0 END) AS active,
                   SUM(CASE WHEN status = 'queued' THEN 1 ELSE 0 END) AS queued,
                   SUM(CASE WHEN status IN ('completed','finished') THEN 1 ELSE 0 END) AS done,
                   SUM(CASE WHEN status = 'failed' THEN 1 ELSE 0 END) AS failed_inst,
                   SUM(COALESCE(completed_tasks,0)) AS ok,
                   SUM(COALESCE(failed_tasks,0)) AS fail
            FROM task_instances
            WHERE queue_id IN ({placeholders})
            GROUP BY queue_id""",
        tuple(queue_ids),
    ).fetchall()
    out = {}
    for r in rows:
        out[r["queue_id"]] = {
            "task_count": r["task_count"] or 0,
            "active_tasks": r["active"] or 0,
            "queued_tasks": r["queued"] or 0,
            "done_instances": r["done"] or 0,
            "failed_instances": r["failed_inst"] or 0,
            "completed_tasks": r["ok"] or 0,
            "failed_tasks": r["fail"] or 0,
        }
    return out


def _queue_status_text(max_concurrent: int, effective: int, st: dict) -> str:
    """队列状态文案：已满 / 运行中 / 空闲。

    与前端 Queues.vue 的 queueStatusText 保持同一口径（按有效并发判断）。
    """
    active = st.get("active_tasks") or 0
    if active >= (effective or 1):
        return "已满"
    return "运行中" if active > 0 else "空闲"


def _sheet_queue_summary(wb, queues: list[dict], stats: dict, title="队列汇总"):
    """一行一队列的队列级汇总（「导出队列列表」的主产物）。"""
    ws = wb.create_sheet(title)
    _write_header(ws, [
        "队列名称", "队列ID", "状态", "任务总数", "运行中", "排队中", "已完成",
        "成功任务数", "失败任务数", "成功率", "并发上限", "有效并发",
        "长尾跳过", "长尾阈值", "描述", "创建人", "创建时间",
    ])
    from .queue_scheduler import effective_concurrent as _eff

    t_count = t_active = t_queued = t_ok = t_fail = 0
    for q in queues:
        st = stats.get(q["id"], {})
        eff = _eff(q.get("max_concurrent"))
        ok = st.get("completed_tasks") or 0
        fail = st.get("failed_tasks") or 0
        t_count += st.get("task_count") or 0
        t_active += st.get("active_tasks") or 0
        t_queued += st.get("queued_tasks") or 0
        t_ok += ok
        t_fail += fail
        ws.append([
            q.get("name") or q.get("id"),
            q.get("id"),
            _queue_status_text(q.get("max_concurrent"), eff, st),
            st.get("task_count") or 0,
            st.get("active_tasks") or 0,
            st.get("queued_tasks") or 0,
            st.get("done_instances") or 0,
            ok,
            fail,
            _rate(ok, fail),
            q.get("max_concurrent") or 1,
            eff,
            "开启" if q.get("tail_skip_enabled") else "关闭",
            f"{q.get('tail_skip_percent') or 10}%",
            q.get("description") or "",
            q.get("created_by") or "",
            _fmt_time(q.get("created_at")),
        ])
    if len(queues) > 1:
        ws.append([
            "合计", "", "", t_count, t_active, t_queued, "", t_ok, t_fail,
            f"{t_ok / (t_ok + t_fail) * 100:.1f}%" if (t_ok + t_fail) else "—",
            "", "", "", "", "", "", "",
        ])
    _fit_columns(ws)
    return ws


def build_queue_summary_workbook(queues: list[dict], stats: dict):
    """「导出队列列表」：只有队列级汇总 + 概览，不含任务明细/配置。"""
    from openpyxl import Workbook
    from datetime import datetime as _dt

    wb = Workbook()
    wb.active.title = "导出概览"
    ov = wb["导出概览"]
    _write_header(ov, ["项目", "值"])
    for k, v in [
        ("导出时间", _dt.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("导出范围", f"{len(queues)} 个队列"),
        ("任务总数", sum((s or {}).get("task_count", 0) for s in stats.values())),
        ("成功任务数合计", sum((s or {}).get("completed_tasks", 0) for s in stats.values())),
        ("失败任务数合计", sum((s or {}).get("failed_tasks", 0) for s in stats.values())),
        ("数据口径", "计数取自 task_instances 与各任务的 complete.jsonl / failed.jsonl 行数"),
        ("说明", "本文件为队列级汇总；需要任务级配置与完整快照请用「导出任务列表」或「导出全部」"),
    ]:
        ov.append([k, v])
    _fit_columns(ov)

    _sheet_queue_summary(wb, queues, stats)
    return wb


def workbook_bytes(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def content_disposition(filename: str) -> str:
    """RFC 5987：filename* 承载中文名，ASCII filename 兜底（同 batch_output）。"""
    safe = (filename or "export.xlsx").replace("/", "_").replace("\\", "_")
    safe = safe.replace('"', "").replace(":", "_")
    ascii_fallback = safe.encode("ascii", "replace").decode("ascii").replace("?", "_")
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(safe)}"
